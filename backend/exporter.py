"""
Module xuất file Excel và CSV chuyên nghiệp:
1. Chế độ "full" (Full Tier): Xuất toàn bộ cây phân cấp Cấu kiện (AS Symbol 'X') & 100% BTP con.
2. Chế độ "missing" (Chỉ chi tiết còn thiếu): Chỉ gom các BTP có con_thieu > 0.
- Tự động tạo hàng Tiêu Đề chuẩn ở dòng 1 và bật sẵn Bộ Lọc (AutoFilter) để người dùng lọc dữ liệu ngay khi mở file.
- Cột Ktra nối được tích hợp vào phần Ghi Chú.
- Hỗ trợ xuất định dạng Excel (.xlsx) đa sheet và CSV (.csv) chuẩn tiếng Việt UTF-8 BOM.
"""

import io
import re
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any, List, Optional

def clean_sheet_title(title: str) -> str:
    """Rút gọn tên sheet để vừa chuẩn Excel (tối đa 31 ký tự)"""
    t = str(title).replace(":", "_").replace("/", "_").replace("\\", "_").replace("?", "").replace("*", "").replace("[", "").replace("]", "")
    return t[:31]

# Định nghĩa các cột chuẩn cho cả Excel và CSV
EXPORT_HEADERS = [
    ("STT", 6, "center"),
    ("Hạng Mục (Sheet)", 16, "left"),
    ("Cấu Kiện Mẹ (Assembly)", 22, "left"),
    ("Tên Cấu Kiện", 20, "left"),
    ("Bản Vẽ (Drawing)", 22, "left"),
    ("Số Lượng Mẹ", 12, "right"),
    ("Tình Trạng Cấu Kiện", 18, "center"),
    ("Mã BTP (Part No)", 18, "left"),
    ("Mã Cắt (Cutting Mark)", 24, "left"),
    ("Tên Chi Tiết (Description)", 24, "left"),
    ("Chủng Loại", 14, "center"),
    ("Quy Cách (Size)", 20, "left"),
    ("Chiều Dài (mm)", 14, "right"),
    ("Vật Liệu", 14, "center"),
    ("Đơn Vị (DVG)", 12, "center"),
    ("Đơn Vị Giao (Tổ)", 16, "left"),
    ("SL Cấu Kiện (Qty/Assy)", 14, "right"),
    ("SL Thiết Kế (Total Qty)", 14, "right"),
    ("Đã Nhận", 12, "right"),
    ("Còn Thiếu", 14, "right"),
    ("Đơn Trọng (kg)", 13, "right"),
    ("KL Thiếu (kg)", 14, "right"),
    ("Kiểm Tra Nối (Ktra Nối)", 14, "center"),
    ("Kế Hoạch Cắt (CP No)", 22, "left"),
    ("Tình Trạng BTP", 25, "left"),
    ("Ghi Chú", 40, "left"),
]

def _collect_export_data(project_data: Dict[str, Any], mode: str, target_sheet: Optional[str]):
    """Gom và lọc dữ liệu cần xuất theo mode và sheet"""
    assemblies = project_data.get("assemblies", [])
    sheets_grouped: Dict[str, List[Dict[str, Any]]] = {}

    for assy in assemblies:
        sheet_name = assy.get("sheet", "Khac")
        
        # Lọc theo hạng mục được chọn nếu có
        if target_sheet and target_sheet.lower() != "all":
            s_up = sheet_name.strip().upper()
            t_up = target_sheet.strip().upper()
            if s_up != t_up and t_up not in s_up and s_up not in t_up:
                continue

        # Xác định danh sách chi tiết cần đưa vào
        if mode == "missing":
            parts_to_include = [p for p in assy.get("parts", []) if (p.get("con_thieu") or 0) > 0]
        else:
            parts_to_include = assy.get("parts", [])

        if parts_to_include or mode == "full":
            if sheet_name not in sheets_grouped:
                sheets_grouped[sheet_name] = []
            sheets_grouped[sheet_name].append({
                "assembly": assy,
                "parts": parts_to_include
            })

    return sheets_grouped

def export_project_excel(
    project_data: Dict[str, Any], 
    mode: str = "missing", 
    target_sheet: Optional[str] = None
) -> io.BytesIO:
    """
    Xuất file Excel chuyên nghiệp:
    - Tiêu đề cột đặt chuẩn ở Dòng 1.
    - Bật sẵn AutoFilter (bộ lọc) và Freeze Panes ở Dòng 1 để người dùng lọc dữ liệu ngay lập tức.
    - Không merge cell giữa các dòng dữ liệu để chức năng lọc / sắp xếp hoạt động chuẩn 100%.
    - Cột Ktra nối được đưa vào phần Ghi Chú.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sheets_grouped = _collect_export_data(project_data, mode, target_sheet)

    if not sheets_grouped:
        ws = wb.create_sheet(title="ThongBao")
        msg = f"Hạng mục {target_sheet} không có chi tiết nào còn thiếu!" if mode == "missing" else "Không có dữ liệu phù hợp."
        ws.cell(1, 1, msg)
        ws.cell(1, 1).font = Font(name="Times New Roman", size=14, bold=True, color="008000")
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Xanh Navy AMECC
    
    font_data = Font(name="Times New Roman", size=10)
    font_missing = Font(name="Times New Roman", size=10, bold=True, color="DC2626")
    fill_missing = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    
    font_completed = Font(name="Times New Roman", size=10, bold=True, color="166534")
    fill_completed = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    
    font_total = Font(name="Times New Roman", size=11, bold=True)
    fill_total = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")

    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    alignments = {
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "left": Alignment(horizontal="left", vertical="center"),
        "right": Alignment(horizontal="right", vertical="center")
    }

    for sheet_name, assy_list in sorted(sheets_grouped.items()):
        ws_title = clean_sheet_title(sheet_name)
        ws = wb.create_sheet(title=ws_title)
        ws.views.sheetView[0].showGridLines = True

        # 1. Dòng 1: Tiêu đề cột chuẩn (Header)
        ws.row_dimensions[1].height = 28
        for col_idx, (h_name, width, align_type) in enumerate(EXPORT_HEADERS, 1):
            cell = ws.cell(1, col_idx, h_name)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = alignments["center"]
            cell.border = border_thin
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # 2. Các dòng dữ liệu (bắt đầu từ Dòng 2)
        cur_row = 2
        stt_counter = 1
        total_tqty_sheet = 0
        total_danhan_sheet = 0
        total_missing_sheet = 0
        total_missing_weight_sheet = 0.0

        for assy_group in assy_list:
            assy = assy_group["assembly"]
            p_list = assy_group["parts"]
            assy_no = assy.get("assembly_no") or assy.get("as_symbol", "")
            assy_name = assy.get("description") or assy.get("as_name", "")
            assy_qty = assy.get("as_qty") or assy.get("assembly_qty") or 1
            dwg = assy.get("dwg", "")
            assy_status = "ĐỦ 100%" if assy.get("status") == "completed" else f"THIẾU ({assy.get('completion_rate', 0)}%)"

            for p in p_list:
                ws.row_dimensions[cur_row].height = 20
                sa = p.get("shape_analysis", {})
                p_qty = p.get("qty") or p.get("qty_per_assy") or 0
                tqty_val = p.get("tqty") or p.get("total_qty") or 0
                da_nhan_val = p.get("da_nhan", 0)
                con_thieu_val = p.get("con_thieu") if p.get("con_thieu") is not None else max(0, tqty_val - da_nhan_val)
                uweight_val = round(float(p.get("uweight") or 0), 2)
                con_thieu_weight = round(float(p.get("con_thieu_weight") or (con_thieu_val * uweight_val)), 1)
                dvg_val = p.get("dvg") or ""
                giao_val = p.get("don_vi_giao") or ""
                desc_val = p.get("description") or p.get("desc") or ""
                chung_loai_val = p.get("chung_loai") or p.get("part_type") or ""
                size_val = p.get("size") or p.get("spec") or ""
                cut_val = p.get("part_cut") or p.get("display_name") or ""
                ktra_noi_val = p.get("ktra_noi") or ""

                total_tqty_sheet += tqty_val
                total_danhan_sheet += da_nhan_val
                total_missing_sheet += con_thieu_val
                total_missing_weight_sheet += con_thieu_weight

                # Format cột còn thiếu
                if con_thieu_val > 0:
                    font_thieu = font_missing
                    fill_thieu = fill_missing
                    if sa.get("has_length_issue"):
                        status_text = "Chưa đủ chiều dài"
                    else:
                        status_text = f"Còn thiếu {con_thieu_val}"
                else:
                    font_thieu = font_completed
                    fill_thieu = fill_completed
                    status_text = "✓ Đã nhận đủ"

                # Ghi chú tổng hợp: bao gồm Ktra nối, vướng thép hình và Remark bản vẽ
                ghi_chu_val = p.get("ghi_chu") or ""
                if not ghi_chu_val:
                    notes = []
                    if ktra_noi_val:
                        notes.append(f"Ktra nối: {ktra_noi_val}")
                    if sa.get("has_length_issue"):
                        notes.append("Chưa đủ chiều dài")
                    if p.get("remark"):
                        notes.append(p["remark"])
                    ghi_chu_val = " | ".join(notes)

                row_vals = [
                    (stt_counter, alignments["center"], font_data, None),
                    (sheet_name, alignments["left"], font_data, None),
                    (assy_no, alignments["left"], font_data, None),
                    (assy_name, alignments["left"], font_data, None),
                    (dwg, alignments["left"], font_data, None),
                    (assy_qty, alignments["right"], font_data, None),
                    (assy_status, alignments["center"], font_data, None),
                    (p.get("part_no", ""), alignments["left"], font_data, None),
                    (cut_val, alignments["left"], font_data, None),
                    (desc_val, alignments["left"], font_data, None),
                    (chung_loai_val, alignments["center"], font_data, None),
                    (size_val, alignments["left"], font_data, None),
                    (p.get("length", ""), alignments["right"], font_data, None),
                    (p.get("material", ""), alignments["center"], font_data, None),
                    (dvg_val, alignments["center"], font_data, None),
                    (giao_val, alignments["left"], font_data, None),
                    (p_qty, alignments["right"], font_data, None),
                    (tqty_val, alignments["right"], font_data, None),
                    (da_nhan_val, alignments["right"], font_data, None),
                    (con_thieu_val, alignments["right"], font_thieu, fill_thieu),
                    (uweight_val if uweight_val > 0 else "-", alignments["right"], font_data, None),
                    (con_thieu_weight if con_thieu_val > 0 else "-", alignments["right"], font_thieu if con_thieu_val > 0 else font_data, fill_thieu if con_thieu_val > 0 else None),
                    (ktra_noi_val, alignments["center"], font_data, None),
                    (sa.get("cutting_no", "") or p.get("cutting_no", "") or "", alignments["left"], font_data, None),
                    (status_text, alignments["left"], font_data, None),
                    (ghi_chu_val, alignments["left"], font_data, None),
                ]

                for c_idx, (val, alignment, font, fill) in enumerate(row_vals, 1):
                    c_cell = ws.cell(cur_row, c_idx, val)
                    c_cell.alignment = alignment
                    c_cell.font = font
                    c_cell.border = border_thin
                    if fill:
                        c_cell.fill = fill

                cur_row += 1
                stt_counter += 1

        # 3. Kích hoạt AutoFilter và Cố định dòng tiêu đề (Freeze Panes)
        last_col = get_column_letter(len(EXPORT_HEADERS))
        last_data_row = max(1, cur_row - 1)
        ws.auto_filter.ref = f"A1:{last_col}{last_data_row}"
        ws.freeze_panes = "A2"

        # 4. Dòng tổng kết ở cuối sheet
        ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=17)
        tot_label_cell = ws.cell(cur_row, 1, "TỔNG CỘNG HẠNG MỤC:")
        tot_label_cell.font = font_total
        tot_label_cell.alignment = Alignment(horizontal="right", vertical="center")
        tot_label_cell.fill = fill_total

        for c in range(1, 18):
            ws.cell(cur_row, c).border = border_thin
            ws.cell(cur_row, c).fill = fill_total

        # Tổng SL Thiết kế (cột 18)
        c_tqty = ws.cell(cur_row, 18, total_tqty_sheet)
        c_tqty.font = font_total
        c_tqty.alignment = alignments["right"]
        c_tqty.fill = fill_total
        c_tqty.border = border_thin

        # Tổng Đã nhận (cột 19)
        c_dn = ws.cell(cur_row, 19, total_danhan_sheet)
        c_dn.font = Font(name="Times New Roman", size=11, bold=True, color="166534")
        c_dn.alignment = alignments["right"]
        c_dn.fill = fill_total
        c_dn.border = border_thin

        # Tổng Còn thiếu (cột 20)
        tot_val_cell = ws.cell(cur_row, 20, total_missing_sheet)
        tot_val_cell.font = Font(name="Times New Roman", size=11, bold=True, color="DC2626")
        tot_val_cell.alignment = alignments["right"]
        tot_val_cell.fill = fill_total
        tot_val_cell.border = border_thin

        # Cột 21 (Đơn Trọng): để trống ở dòng tổng
        c_u = ws.cell(cur_row, 21, "-")
        c_u.font = font_total
        c_u.alignment = alignments["center"]
        c_u.fill = fill_total
        c_u.border = border_thin

        # Cột 22 (Tổng KL Thiếu):
        c_w = ws.cell(cur_row, 22, round(total_missing_weight_sheet, 1))
        c_w.font = Font(name="Times New Roman", size=11, bold=True, color="DC2626")
        c_w.alignment = alignments["right"]
        c_w.fill = fill_total
        c_w.border = border_thin

        for c in range(23, len(EXPORT_HEADERS) + 1):
            ws.cell(cur_row, c).border = border_thin
            ws.cell(cur_row, c).fill = fill_total

        ws.row_dimensions[cur_row].height = 24

        # Tự động căn chỉnh độ rộng cột vừa vặn với nội dung (tối đa 50 ký tự)
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row == cur_row:  # Bỏ qua dòng tổng cộng đã merge
                    continue
                v_str = str(cell.value or "")
                if len(v_str) > max_len:
                    max_len = len(v_str)
            if max_len > 0:
                current_w = ws.column_dimensions[col_letter].width or 12
                ws.column_dimensions[col_letter].width = min(max(current_w, max_len + 3), 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_project_csv(
    project_data: Dict[str, Any], 
    mode: str = "missing", 
    target_sheet: Optional[str] = None
) -> io.BytesIO:
    """
    Xuất file CSV chuẩn UTF-8 BOM:
    - Mở trực tiếp bằng Microsoft Excel không bị lỗi font tiếng Việt.
    - Dòng 1 là tiêu đề cột để người dùng bật Filter (Ctrl+Shift+L) lọc cực dễ.
    - Cột Ktra nối được đưa vào phần Ghi Chú.
    """
    sheets_grouped = _collect_export_data(project_data, mode, target_sheet)

    # Sử dụng StringIO trước để ghi CSV
    string_output = io.StringIO()
    writer = csv.writer(string_output, quoting=csv.QUOTE_MINIMAL)

    # 1. Ghi dòng tiêu đề
    headers = [h[0] for h in EXPORT_HEADERS]
    writer.writerow(headers)

    stt_counter = 1
    for sheet_name, assy_list in sorted(sheets_grouped.items()):
        for assy_group in assy_list:
            assy = assy_group["assembly"]
            p_list = assy_group["parts"]
            assy_no = assy.get("assembly_no") or assy.get("as_symbol", "")
            assy_name = assy.get("description") or assy.get("as_name", "")
            assy_qty = assy.get("as_qty") or assy.get("assembly_qty") or 1
            dwg = assy.get("dwg", "")
            assy_status = "ĐỦ 100%" if assy.get("status") == "completed" else f"THIẾU ({assy.get('completion_rate', 0)}%)"

            for p in p_list:
                sa = p.get("shape_analysis", {})
                p_qty = p.get("qty") or p.get("qty_per_assy") or 0
                tqty_val = p.get("tqty") or p.get("total_qty") or 0
                da_nhan_val = p.get("da_nhan", 0)
                con_thieu_val = p.get("con_thieu") if p.get("con_thieu") is not None else max(0, tqty_val - da_nhan_val)

                if con_thieu_val > 0:
                    status_text = "Chưa đủ chiều dài" if sa.get("has_length_issue") else f"Còn thiếu {con_thieu_val}"
                else:
                    status_text = "Đã nhận đủ"

                ktra_noi_val = p.get("ktra_noi") or ""
                ghi_chu_val = p.get("ghi_chu") or ""
                if not ghi_chu_val:
                    notes = []
                    if ktra_noi_val:
                        notes.append(f"Ktra nối: {ktra_noi_val}")
                    if sa.get("has_length_issue"):
                        notes.append("Chưa đủ chiều dài")
                    if p.get("remark"):
                        notes.append(p["remark"])
                    ghi_chu_val = " | ".join(notes)

                uweight_val = round(float(p.get("uweight") or 0), 2)
                con_thieu_weight = round(float(p.get("con_thieu_weight") or (con_thieu_val * uweight_val)), 1)
                dvg_val = p.get("dvg") or ""
                giao_val = p.get("don_vi_giao") or ""
                desc_val = p.get("description") or p.get("desc") or ""
                chung_loai_val = p.get("chung_loai") or p.get("part_type") or ""
                size_val = p.get("size") or p.get("spec") or ""
                cut_val = p.get("part_cut") or p.get("display_name") or ""

                row = [
                    stt_counter,
                    sheet_name,
                    assy_no,
                    assy_name,
                    dwg,
                    assy_qty,
                    assy_status,
                    p.get("part_no", ""),
                    cut_val,
                    desc_val,
                    chung_loai_val,
                    size_val,
                    p.get("length", ""),
                    p.get("material", ""),
                    dvg_val,
                    giao_val,
                    p_qty,
                    tqty_val,
                    da_nhan_val,
                    con_thieu_val,
                    uweight_val if uweight_val > 0 else "",
                    con_thieu_weight if con_thieu_val > 0 else "",
                    ktra_noi_val,
                    sa.get("cutting_no", "") or p.get("cutting_no", "") or "",
                    status_text,
                    ghi_chu_val
                ]
                writer.writerow(row)
                stt_counter += 1

    # Chuyển đổi sang bytes với mã hóa UTF-8-SIG (Excel mở chuẩn dấu tiếng Việt)
    csv_bytes = string_output.getvalue().encode('utf-8-sig')
    return io.BytesIO(csv_bytes)

def normalize_size_exporter(size_str: Any) -> str:
    """Chuẩn hóa quy cách (Size) thành mã kích thước/độ dày chính (ví dụ -PL 75*1500 -> PL75)"""
    if not size_str:
        return "Khác"
    s = str(size_str).strip()
    s = re.sub(r'^[-_]\s*', '', s)
    pl_match = re.match(r'^PL\s*[-_]?\s*(\d+)', s, re.IGNORECASE)
    if pl_match:
        return f"PL{pl_match.group(1)}"
    shape_match = re.match(r'^([A-Za-z]+)\s*[-_]?\s*(\d+)', s)
    if shape_match:
        pref = shape_match.group(1).upper()
        if pref in ["H", "I", "L", "C", "U", "V", "T", "PIPE", "TUBE", "FB"]:
            return f"{pref}{shape_match.group(2)}"
    parts = re.split(r'[*xX\s]', s)
    return parts[0].strip().upper() if parts[0].strip() else s

def export_size_summary_excel(
    project_data: Dict[str, Any],
    target_sheet: Optional[str] = None
) -> io.BytesIO:
    """
    Xuất Bảng Tổng Hợp Quy Cách (SIZE) BOM chuyên nghiệp:
    - Tổng hợp theo độ dày/kích thước (như PL75, PL10, H300, I200...)
    - Tính tổng SL thiết kế, đã nhận, còn thiếu, khối lượng thiếu, tỷ lệ %
    - Tiêu đề Xanh Navy AMECC, AutoFilter, Freeze Panes, dòng tổng cộng công thức
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TongHop_SIZE"
    ws.views.sheetView[0].showGridLines = True

    assemblies = project_data.get("assemblies", [])
    proj_code = project_data.get("project_id", "DuAn")

    # Gom nhóm theo Size
    size_map: Dict[str, Dict[str, Any]] = {}

    for assy in assemblies:
        sheet_name = assy.get("sheet", "Khac")
        if target_sheet and target_sheet.lower() != "all":
            s_up = sheet_name.strip().upper()
            t_up = target_sheet.strip().upper()
            if s_up != t_up and t_up not in s_up and s_up not in t_up:
                continue

        for p in assy.get("parts", []):
            raw_sz = p.get("size") or p.get("spec") or ""
            norm_sz = normalize_size_exporter(raw_sz)
            chung_loai = p.get("chung_loai") or p.get("part_type") or ("Thép tấm" if norm_sz.startswith("PL") else "Thép hình")
            
            tqty = p.get("tqty") or p.get("total_qty") or 0
            da_nhan = p.get("da_nhan") or 0
            con_thieu = p.get("con_thieu") if p.get("con_thieu") is not None else max(0, tqty - da_nhan)
            uweight = float(p.get("uweight") or 0.0)
            con_thieu_w = float(p.get("con_thieu_weight") or (con_thieu * uweight))

            if norm_sz not in size_map:
                size_map[norm_sz] = {
                    "size": norm_sz,
                    "chung_loai": chung_loai,
                    "tqty": 0,
                    "da_nhan": 0,
                    "con_thieu": 0,
                    "con_thieu_weight": 0.0,
                    "uweights": [],
                    "part_count": 0
                }

            size_map[norm_sz]["tqty"] += tqty
            size_map[norm_sz]["da_nhan"] += da_nhan
            size_map[norm_sz]["con_thieu"] += con_thieu
            size_map[norm_sz]["con_thieu_weight"] += con_thieu_w
            if uweight > 0:
                size_map[norm_sz]["uweights"].append(uweight)
            size_map[norm_sz]["part_count"] += 1

    headers = [
        ("STT", 6, "center"),
        ("Chủng Loại", 16, "center"),
        ("Quy Cách (SIZE)", 20, "left"),
        ("Số Chi Tiết", 12, "right"),
        ("Tổng SL Thiết Kế", 18, "right"),
        ("Tổng Đã Nhận", 16, "right"),
        ("Còn Thiếu", 16, "right"),
        ("Đơn Trọng TB (kg)", 18, "right"),
        ("Tổng KL Thiếu (kg)", 20, "right"),
        ("Tỷ Lệ Nhận (%)", 16, "right"),
        ("Trạng Thái", 18, "center")
    ]

    font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1A3C8F", end_color="1A3C8F", fill_type="solid") # AMECC Blue
    
    font_data = Font(name="Times New Roman", size=10)
    font_missing = Font(name="Times New Roman", size=10, bold=True, color="DC2626")
    fill_missing = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    
    font_completed = Font(name="Times New Roman", size=10, bold=True, color="166534")
    fill_completed = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    
    font_total = Font(name="Times New Roman", size=11, bold=True)
    fill_total = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")

    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    alignments = {
        "center": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center"),
        "right": Alignment(horizontal="right", vertical="center")
    }

    # Dòng 1: Header
    ws.row_dimensions[1].height = 28
    for col_idx, (h_name, width, align_type) in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, h_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = alignments["center"]
        cell.border = border_thin
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Sắp xếp: Ưu tiên còn thiếu nhiều nhất lên đầu
    sorted_sizes = sorted(
        size_map.values(),
        key=lambda x: (x["con_thieu"] <= 0, -x["con_thieu"], x["size"])
    )

    cur_row = 2
    for stt, sdata in enumerate(sorted_sizes, 1):
        ws.row_dimensions[cur_row].height = 20
        tqty = sdata["tqty"]
        da_nhan = sdata["da_nhan"]
        con_thieu = sdata["con_thieu"]
        pct = round((da_nhan / tqty * 100), 1) if tqty > 0 else 0.0
        avg_uw = round(sum(sdata["uweights"]) / len(sdata["uweights"]), 2) if sdata["uweights"] else 0.0
        missing_w = round(sdata["con_thieu_weight"], 1)

        if con_thieu == 0 and tqty > 0:
            status_text = "ĐỦ 100%"
            f_color = font_completed
            f_fill = fill_completed
        elif da_nhan == 0:
            status_text = "CHƯA NHẬN"
            f_color = font_missing
            f_fill = fill_missing
        else:
            status_text = f"THIẾU {con_thieu}"
            f_color = font_missing
            f_fill = fill_missing

        row_vals = [
            (stt, alignments["center"], font_data, None, None),
            (sdata["chung_loai"], alignments["center"], font_data, None, None),
            (sdata["size"], alignments["left"], Font(name="Times New Roman", size=10, bold=True), None, None),
            (sdata["part_count"], alignments["right"], font_data, None, "#,##0"),
            (tqty, alignments["right"], font_data, None, "#,##0"),
            (da_nhan, alignments["right"], font_completed if da_nhan > 0 else font_data, None, "#,##0"),
            (con_thieu, alignments["right"], f_color, f_fill if con_thieu > 0 else None, "#,##0"),
            (avg_uw if avg_uw > 0 else "-", alignments["right"], font_data, None, "#,##0.00" if avg_uw > 0 else None),
            (missing_w if missing_w > 0 else "-", alignments["right"], f_color if missing_w > 0 else font_data, f_fill if missing_w > 0 else None, "#,##0.0" if missing_w > 0 else None),
            (f"{pct}%", alignments["right"], f_color, None, None),
            (status_text, alignments["center"], f_color, f_fill, None)
        ]

        for c_idx, (val, alignment, font, fill, num_format) in enumerate(row_vals, 1):
            c_cell = ws.cell(cur_row, c_idx, val)
            c_cell.alignment = alignment
            c_cell.font = font
            c_cell.border = border_thin
            if fill:
                c_cell.fill = fill
            if num_format:
                c_cell.number_format = num_format

        cur_row += 1

    # Freeze panes & AutoFilter
    last_col = get_column_letter(len(headers))
    last_data_row = max(1, cur_row - 1)
    ws.auto_filter.ref = f"A1:{last_col}{last_data_row}"
    ws.freeze_panes = "A2"

    # Dòng tổng cộng
    ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=3)
    tot_label = ws.cell(cur_row, 1, "TỔNG CỘNG SIZE:")
    tot_label.font = font_total
    tot_label.alignment = Alignment(horizontal="right", vertical="center")
    tot_label.fill = fill_total

    for c in range(1, 4):
        ws.cell(cur_row, c).border = border_thin
        ws.cell(cur_row, c).fill = fill_total

    # Công thức Subtotal cho Part count (Col 4), Tqty (Col 5), Đã nhận (Col 6), Còn thiếu (Col 7), KL Thiếu (Col 9)
    subtotal_map = [
        (4, "#,##0"),
        (5, "#,##0"),
        (6, "#,##0"),
        (7, "#,##0"),
        (9, "#,##0.0")
    ]
    for col_c, num_fmt in subtotal_map:
        col_let = get_column_letter(col_c)
        c_cell = ws.cell(cur_row, col_c)
        c_cell.value = f"=SUBTOTAL(9, {col_let}2:{col_let}{last_data_row})"
        c_cell.font = font_total
        c_cell.alignment = alignments["right"]
        c_cell.fill = fill_total
        c_cell.border = border_thin
        c_cell.number_format = num_fmt

    # Cột 8 (Đơn trọng TB) để trống ở dòng tổng
    c_empty8 = ws.cell(cur_row, 8, "-")
    c_empty8.font = font_total
    c_empty8.alignment = alignments["center"]
    c_empty8.fill = fill_total
    c_empty8.border = border_thin

    # Cột 10 & 11
    for c in [10, 11]:
        c_cell = ws.cell(cur_row, c, "")
        c_cell.fill = fill_total
        c_cell.border = border_thin

    ws.row_dimensions[cur_row].height = 24

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_daily_receipts_excel(
    project_data: Dict[str, Any],
    date_str: str,
    target_sheet: Optional[str] = None
) -> io.BytesIO:
    """
    Xuất Danh Sách BTP Nhận Theo Ngày:
    - Lọc các chi tiết BTP có lịch sử nhận trong ngày date_str
    - Phân tách theo Hạng Mục, Cấu Kiện Mẹ, Mã Cắt, Kích Thước, Số Lượng, Khối Lượng
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    clean_date_title = date_str.replace("/", "-")
    ws.title = clean_sheet_title(f"Nhan_{clean_date_title}")
    ws.views.sheetView[0].showGridLines = True

    assemblies = project_data.get("assemblies", [])
    proj_code = project_data.get("project_id", "DuAn")

    headers = [
        ("STT", 6, "center"),
        ("Hạng Mục (Sheet)", 16, "left"),
        ("Cấu Kiện Mẹ", 20, "left"),
        ("Tên Cấu Kiện", 22, "left"),
        ("Mã BTP (Part No)", 18, "left"),
        ("Mã Cắt", 20, "left"),
        ("Quy Cách (Size)", 18, "left"),
        ("Vật Liệu", 14, "center"),
        ("Đơn Vị Giao (Tổ)", 16, "left"),
        ("Ngày Nhận", 14, "center"),
        ("SL Nhận Trong Ngày", 18, "right"),
        ("Đơn Trọng (kg)", 14, "right"),
        ("KL Nhận (kg)", 16, "right"),
        ("Kiểm Tra Nối", 14, "center"),
        ("Ghi Chú", 30, "left")
    ]

    font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1A3C8F", end_color="1A3C8F", fill_type="solid")
    font_data = Font(name="Times New Roman", size=10)
    font_total = Font(name="Times New Roman", size=11, bold=True)
    fill_total = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")

    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    alignments = {
        "center": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center"),
        "right": Alignment(horizontal="right", vertical="center")
    }

    # Dòng 1: Header
    ws.row_dimensions[1].height = 28
    for col_idx, (h_name, width, align_type) in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, h_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = alignments["center"]
        cell.border = border_thin
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    cur_row = 2
    stt = 1
    total_qty_day = 0
    total_weight_day = 0.0

    target_d_clean = date_str.strip()

    for assy in assemblies:
        sheet_name = assy.get("sheet", "Khac")
        if target_sheet and target_sheet.lower() != "all":
            s_up = sheet_name.strip().upper()
            t_up = target_sheet.strip().upper()
            if s_up != t_up and t_up not in s_up and s_up not in t_up:
                continue

        assy_no = assy.get("assembly_no") or assy.get("as_symbol", "")
        assy_name = assy.get("description") or assy.get("as_name", "")

        for p in assy.get("parts", []):
            dates_rec = p.get("dates_received") or {}
            # Khớp ngày linh hoạt
            matched_qty = 0
            for d_k, d_v in dates_rec.items():
                if d_k.strip() == target_d_clean or target_d_clean in d_k:
                    matched_qty += (d_v or 0)

            if matched_qty > 0:
                ws.row_dimensions[cur_row].height = 20
                uweight = float(p.get("uweight") or 0.0)
                rec_weight = round(matched_qty * uweight, 1) if uweight else 0.0

                total_qty_day += matched_qty
                total_weight_day += rec_weight

                row_vals = [
                    (stt, alignments["center"], font_data, None, None),
                    (sheet_name, alignments["left"], font_data, None, None),
                    (assy_no, alignments["left"], font_data, None, None),
                    (assy_name, alignments["left"], font_data, None, None),
                    (p.get("part_no", ""), alignments["left"], font_data, None, None),
                    (p.get("part_cut", "") or p.get("display_name", ""), alignments["left"], font_data, None, None),
                    (p.get("size", "") or p.get("spec", ""), alignments["left"], font_data, None, None),
                    (p.get("material", ""), alignments["center"], font_data, None, None),
                    (p.get("don_vi_giao", "") or p.get("dvg", ""), alignments["left"], font_data, None, None),
                    (target_d_clean, alignments["center"], font_data, None, None),
                    (matched_qty, alignments["right"], Font(name="Times New Roman", size=10, bold=True, color="166534"), None, "#,##0"),
                    (uweight if uweight > 0 else "-", alignments["right"], font_data, None, "#,##0.00" if uweight > 0 else None),
                    (rec_weight if rec_weight > 0 else "-", alignments["right"], Font(name="Times New Roman", size=10, bold=True, color="166534"), None, "#,##0.0" if rec_weight > 0 else None),
                    (p.get("ktra_noi", ""), alignments["center"], font_data, None, None),
                    (p.get("ghi_chu", "") or p.get("remark", ""), alignments["left"], font_data, None, None)
                ]

                for c_idx, (val, alignment, font, fill, num_format) in enumerate(row_vals, 1):
                    c_cell = ws.cell(cur_row, c_idx, val)
                    c_cell.alignment = alignment
                    c_cell.font = font
                    c_cell.border = border_thin
                    if fill:
                        c_cell.fill = fill
                    if num_format:
                        c_cell.number_format = num_format

                cur_row += 1
                stt += 1

    last_col = get_column_letter(len(headers))
    last_data_row = max(1, cur_row - 1)
    ws.auto_filter.ref = f"A1:{last_col}{last_data_row}"
    ws.freeze_panes = "A2"

    # Dòng tổng cộng
    ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=10)
    tot_label = ws.cell(cur_row, 1, f"TỔNG CỘNG NGÀY {target_d_clean}:")
    tot_label.font = font_total
    tot_label.alignment = Alignment(horizontal="right", vertical="center")
    tot_label.fill = fill_total

    for c in range(1, 11):
        ws.cell(cur_row, c).border = border_thin
        ws.cell(cur_row, c).fill = fill_total

    c_qty = ws.cell(cur_row, 11)
    c_qty.value = f"=SUBTOTAL(9, K2:K{last_data_row})"
    c_qty.font = Font(name="Times New Roman", size=11, bold=True, color="166534")
    c_qty.alignment = alignments["right"]
    c_qty.fill = fill_total
    c_qty.border = border_thin
    c_qty.number_format = "#,##0"

    c_uw = ws.cell(cur_row, 12, "-")
    c_uw.font = font_total
    c_uw.alignment = alignments["center"]
    c_uw.fill = fill_total
    c_uw.border = border_thin

    c_w = ws.cell(cur_row, 13)
    c_w.value = f"=SUBTOTAL(9, M2:M{last_data_row})"
    c_w.font = Font(name="Times New Roman", size=11, bold=True, color="166534")
    c_w.alignment = alignments["right"]
    c_w.fill = fill_total
    c_w.border = border_thin
    c_w.number_format = "#,##0.0"

    for c in range(14, len(headers) + 1):
        c_cell = ws.cell(cur_row, c, "")
        c_cell.fill = fill_total
        c_cell.border = border_thin

    ws.row_dimensions[cur_row].height = 24

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# Tương thích ngược
export_missing_parts_excel = export_project_excel

